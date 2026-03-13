import json
import os
import datetime
import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from ortools.sat.python import cp_model

MAX_SOLVE_TIME = 450.0

BASE_DIR = r"C:\Users\psobo\PycharmProjects\clean_VKR-feature-synthetic-data-mode\ortools"
DRIVERS_FILE = os.path.join(BASE_DIR, "drivers_9.json")  # Сюда кладешь водителей всех маршрутов
SCHEDULE_FILE = os.path.join(BASE_DIR, "schedule_9.json")  # Сюда кладешь расписания всех маршрутов

OUTPUT_DRIVERS_EXCEL = os.path.join(BASE_DIR, "report_drivers_ortools.xlsx")
OUTPUT_SCHEDULE_BOOK = os.path.join(BASE_DIR, "report_schedule_book.xlsx")

MONTH_MAP = {
    "Январь": 1, "Февраль": 2, "Март": 3, "Апрель": 4, "Май": 5, "Июнь": 6,
    "Июль": 7, "Август": 8, "Сентябрь": 9, "Октябрь": 10, "Ноябрь": 11, "Декабрь": 12
}
DAYS_RU = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]


def parse_time(time_str):
    h, m = map(int, time_str.split(':'))
    return h * 60 + m


def format_mins_for_excel(minutes):
    if minutes is None: return "-"
    h = int(minutes // 60)
    m = int(minutes % 60)
    return f"{h}ч {m:02d}м"


def get_clock_time_str(mins_from_midnight):
    m = mins_from_midnight % 1440
    return f"{int(m // 60):02d}:{int(m % 60):02d}"


def get_day_of_week_str(year, month, day):
    try:
        return DAYS_RU[datetime.date(year, month, day).weekday()]
    except:
        return "?"


def load_data():
    with open(SCHEDULE_FILE, 'r', encoding='utf-8') as f:
        schedule_data = json.load(f)

    shifts_info = {"рабочий": {}, "выходной": {}}

    for sched in schedule_data:
        day_type = sched.get("день")
        route_num = str(sched.get("маршрут", "Unknown"))

        for t in sched.get("трамваи", []):
            t_num = f"{route_num}_{t['номер']}"

            if t_num not in shifts_info[day_type]:
                shifts_info[day_type][t_num] = {}

            for s_num in [1, 2]:
                s_key = f"смена_{s_num}"
                shift_data = t.get(s_key)

                if shift_data and isinstance(shift_data, dict):
                    if "отправление" in shift_data and "прибытие" in shift_data:
                        st = parse_time(shift_data["отправление"])
                        en = parse_time(shift_data["прибытие"])
                        if en < st: en += 24 * 60
                        shifts_info[day_type][t_num][s_num] = {"start": st, "end": en, "duration": en - st}

    with open(DRIVERS_FILE, 'r', encoding='utf-8') as f:
        drivers_raw = json.load(f)

    year = int(drivers_raw.get("year", 2026))
    month = MONTH_MAP.get(drivers_raw.get("month", "Январь"), 1)

    drivers_ids = []
    availabilities = {}
    driver_schedules = {}

    def get_day_type(day):
        try:
            return "выходной" if datetime.date(year, month, day).weekday() >= 5 else "рабочий"
        except ValueError:
            return None

    for d in drivers_raw.get("drivers", []):
        d_id = str(d.get("tab_number") or d.get("id") or "unknown")
        drivers_ids.append(d_id)
        driver_schedules[d_id] = d.get("schedule") or d.get("schedule_pattern") or d.get("mode") or "4x2"

        availabilities[d_id] = {day: 0 for day in range(1, 32)}
        for day_record in d.get("days", []):
            day = day_record.get("day")
            val = str(day_record.get("value")).upper().strip()
            if val == "1":
                availabilities[d_id][day] = 1
            elif val == "2":
                availabilities[d_id][day] = 2
            elif val in ['Р', 'Я', 'S', 'WORK']:
                availabilities[d_id][day] = -1

    return drivers_ids, shifts_info, availabilities, get_day_type, driver_schedules, year, month


def export_drivers_report(drivers, used_drivers, driver_timeline, shifts_info, availabilities, get_day_type,
                          driver_schedules, num_days=31):
    rows = []
    header = ["Таб.№", "График", "Смен", "Часов"] + [f"{d}" for d in range(1, num_days + 1)]
    sorted_used_drivers = sorted(list(used_drivers), key=lambda x: int(x) if x.isdigit() else x)

    for d in sorted_used_drivers:
        shifts_list = driver_timeline[d]
        total_shifts = len(shifts_list)
        total_hours = sum(s["duration"] for s in shifts_list) / 60.0
        row = {"Таб.№": d, "График": driver_schedules.get(d, "?"), "Смен": total_shifts, "Часов": round(total_hours, 1)}

        for day in range(1, num_days + 1):
            day_type = get_day_type(day)
            if not day_type:
                row[str(day)] = ""
                continue

            worked_shift = next((s for s in shifts_list if s["day"] == day), None)
            if worked_shift:
                work_time = format_mins_for_excel(worked_shift["duration"])
                idx = shifts_list.index(worked_shift)
                rest_before = shifts_list[idx - 1]["rest_after"] if idx > 0 else None
                row[str(day)] = f"{work_time} (отд {format_mins_for_excel(rest_before)})"
            else:
                plan = availabilities[d].get(day, 0)
                if plan == 0:
                    row[str(day)] = "В"
                else:
                    past_shifts = [s for s in shifts_list if s["day"] < day]
                    if not past_shifts:
                        row[str(day)] = "РЕЗЕРВ"
                    else:
                        last_s = past_shifts[-1]
                        hypo_start = (day - 1) * 1440 + (300 if plan == 1 else 840)
                        planned_weekends = sum(
                            1 for d_idx in range(last_s["day"] + 1, day) if availabilities[d].get(d_idx, 0) == 0)
                        if (hypo_start - last_s["abs_end"]) < (last_s["duration"] * 2) + (
                        2520 if planned_weekends > 0 else 0):
                            row[str(day)] = "МАЛО ОТДЫХА"
                        else:
                            row[str(day)] = "РЕЗЕРВ"
        rows.append(row)

    df = pd.DataFrame(rows, columns=header)
    stat_total = {"Таб.№": f"Всего: {len(used_drivers)} чел.", "График": "", "Смен": "", "Часов": "ВСЕГО СМЕН:"}
    stat_unclosed = {"Таб.№": "", "График": "", "Смен": "", "Часов": "НЕЗАКРЫТЫХ СМЕН:"}

    for day in range(1, num_days + 1):
        day_type = get_day_type(day)
        if not day_type: continue
        total_day_shifts = sum(len(shifts) for shifts in shifts_info[day_type].values())
        closed_shifts = sum(1 for d in used_drivers for s in driver_timeline[d] if s["day"] == day)
        stat_total[str(day)] = total_day_shifts
        stat_unclosed[str(day)] = total_day_shifts - closed_shifts

    df_final = pd.concat([df, pd.DataFrame([stat_total, stat_unclosed], columns=header)], ignore_index=True)
    writer = pd.ExcelWriter(OUTPUT_DRIVERS_EXCEL, engine='openpyxl')
    df_final.to_excel(writer, index=False, sheet_name="Расписание")

    ws = writer.sheets["Расписание"]
    ws.freeze_panes = 'E2'

    fill_work = PatternFill("solid", fgColor="C6EFCE")
    fill_rest = PatternFill("solid", fgColor="F2F2F2")
    fill_reserve = PatternFill("solid", fgColor="FFC7CE")
    fill_warn = PatternFill("solid", fgColor="FF9900")
    fill_stat = PatternFill("solid", fgColor="D9E1F2")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'),
                    bottom=Side(style='thin'))

    for r_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        is_stat_row = r_idx > len(used_drivers) + 1
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            val = str(cell.value) if cell.value else ""
            if is_stat_row:
                cell.fill = fill_stat;
                cell.font = Font(bold=True)
                if "НЕЗАКРЫТЫХ" in val or (
                        cell.column > 4 and cell.value and int(cell.value) > 0 and row[3].value == "НЕЗАКРЫТЫХ СМЕН:"):
                    cell.font = Font(bold=True, color="FF0000")
                continue
            if "(отд" in val:
                cell.fill = fill_work
            elif "В" == val:
                cell.fill = fill_rest
            elif "РЕЗЕРВ" in val:
                cell.fill = fill_reserve
            elif "МАЛО ОТДЫХА" in val:
                cell.fill = fill_warn; cell.font = Font(bold=True)

    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = min(
        max((len(str(c.value)) for c in col if c.value), default=10) + 2, 25)
    writer.close()


def export_schedule_book(driver_timeline, shifts_info, get_day_type, year, month, num_days=31):
    shift_assignments = {}
    for d, shifts in driver_timeline.items():
        for s in shifts:
            shift_assignments[(s["day"], s["tram"], s["shift"])] = d

    rows = []
    for day in range(1, num_days + 1):
        day_type = get_day_type(day)
        if not day_type: continue

        dow = get_day_of_week_str(year, month, day)
        date_str = f"{day:02d}.{month:02d} ({dow})"

        def sort_key(tram_id):
            r_str, t_str = tram_id.split('_', 1)
            r_int = int(r_str) if r_str.isdigit() else 9999
            t_int = int(t_str) if t_str.isdigit() else 9999
            return (r_int, t_int)

        trams_for_day = sorted(list(shifts_info[day_type].keys()), key=sort_key)

        for t in trams_for_day:
            route_str, tram_str = t.split('_', 1)

            row_data = {
                "Дата": date_str,
                "Маршрут": route_str,  # НОВАЯ КОЛОНКА
                "Вагон": int(tram_str) if tram_str.isdigit() else tram_str,
                "I Смена (Водитель)": "", "I Время": "",
                "II Смена (Водитель)": "", "II Время": "",
                "Проблемы": "", "_day_int": day
            }
            issues = []

            for s_num in [1, 2]:
                if s_num in shifts_info[day_type][t]:
                    s_info = shifts_info[day_type][t][s_num]
                    time_str = f"{get_clock_time_str(s_info['start'])} - {get_clock_time_str(s_info['end'])}"
                    driver = shift_assignments.get((day, t, s_num))

                    if s_num == 1:
                        row_data["I Время"] = time_str
                        if driver:
                            row_data["I Смена (Водитель)"] = driver
                        else:
                            row_data["I Смена (Водитель)"] = "[FAIL]"; issues.append("Нет водителя (1)")
                    else:
                        row_data["II Время"] = time_str
                        if driver:
                            row_data["II Смена (Водитель)"] = driver
                        else:
                            row_data["II Смена (Водитель)"] = "[FAIL]"; issues.append("Нет водителя (2)")

            if issues: row_data["Проблемы"] = " ".join(issues)
            if row_data["I Время"] or row_data["II Время"]: rows.append(row_data)

    df = pd.DataFrame(rows).drop(columns=["_day_int"])
    writer = pd.ExcelWriter(OUTPUT_SCHEDULE_BOOK, engine='openpyxl')
    df.to_excel(writer, index=False, sheet_name="Журнал", startrow=5)

    ws = writer.sheets["Журнал"]
    fill_success = PatternFill("solid", fgColor="C6EFCE")
    fill_fail = PatternFill("solid", fgColor="FF5555")
    fill_header = PatternFill("solid", fgColor="4F81BD")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'),
                    bottom=Side(style='thin'))
    thick_bottom = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'),
                          bottom=Side(style='thick'))

    ws['A1'] = "ЛЕГЕНДА:"
    ws['A2'] = "Смена закрыта (OK)";
    ws['A2'].fill = fill_success
    ws['B2'] = "Срыв (Нет водителя)";
    ws['B2'].fill = fill_fail

    for cell in ws[6]:
        cell.font = Font(bold=True, color="FFFFFF");
        cell.fill = fill_header;
        cell.alignment = Alignment(horizontal='center')

    for r_idx, row in enumerate(ws.iter_rows(min_row=7, max_row=ws.max_row), start=7):
        is_day_end = False
        if r_idx < ws.max_row:
            next_date = ws.cell(row=r_idx + 1, column=1).value
            if next_date != row[0].value: is_day_end = True
        else:
            is_day_end = True

        for cell in row:
            cell.border = thick_bottom if is_day_end else border
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for c_idx in [3, 5]:
            val = str(row[c_idx].value) if row[c_idx].value else ""
            if "[FAIL]" in val:
                row[c_idx].fill = fill_fail
                row[c_idx].value = ""
            elif val.strip():
                row[c_idx].fill = fill_success

    ws.freeze_panes = 'A7'
    ws.column_dimensions['A'].width = 15  # Дата
    ws.column_dimensions['B'].width = 12  # Маршрут
    ws.column_dimensions['C'].width = 10  # Вагон
    ws.column_dimensions['D'].width = 20  # I Водитель
    ws.column_dimensions['E'].width = 15  # I Время
    ws.column_dimensions['F'].width = 20  # II Водитель
    ws.column_dimensions['G'].width = 15  # II Время
    ws.column_dimensions['H'].width = 30  # Проблемы
    writer.close()
    print(f"-> Журнал нарядов: {OUTPUT_SCHEDULE_BOOK}")


def precheck_feasibility(drivers, shifts_info, availabilities, get_day_type, num_days):

    print("\nПРЕДВАРИТЕЛЬНАЯ ПРОВЕРКА ДАННЫХ")
    all_good = True

    for day in range(1, num_days + 1):
        day_type = get_day_type(day)
        if not day_type: continue

        s1_needed = sum(1 for t in shifts_info[day_type].values() if 1 in t)
        s2_needed = sum(1 for t in shifts_info[day_type].values() if 2 in t)

        s1_avail, s2_avail, any_avail = 0, 0, 0
        for d in drivers:
            plan = availabilities[d].get(day, 0)
            if plan == 1:
                s1_avail += 1
            elif plan == 2:
                s2_avail += 1
            elif plan == -1:
                any_avail += 1

        if (s1_avail + any_avail < s1_needed) or (s2_avail + any_avail < s2_needed) or (
                s1_avail + s2_avail + any_avail < s1_needed + s2_needed):
            print(f"  [!] КРИТИЧЕСКАЯ НЕХВАТКА ЛЮДЕЙ: ДЕНЬ {day:02d} ({day_type})")
            print(f"      Смена 1 (Утро) : Нужно {s1_needed}, Доступно {s1_avail} (+ {any_avail} универс.)")
            print(f"      Смена 2 (Вечер): Нужно {s2_needed}, Доступно {s2_avail} (+ {any_avail} универс.)")
            all_good = False

    if not all_good:
        print("ОШИБКА: Математически невозможно закрыть все смены без дыр.")
        print("Алгоритм остановлен. Добавить водителей в базу")
        return False

    print("Проверка пройдена! Людей достаточно для старта.")
    return True


def solve_schedule():
    drivers, shifts_info, availabilities, get_day_type, driver_schedules, year, month = load_data()
    num_days = 31

    if not precheck_feasibility(drivers, shifts_info, availabilities, get_day_type, num_days):
        return

    print(f"\nСборка модели OR-Tools")
    model = cp_model.CpModel()
    work = {}

    print("Шаг 1/5: Генерация разреженной матрицы")
    for d in drivers:
        for day in range(1, num_days + 1):
            allowed_shift = availabilities[d].get(day, 0)
            if allowed_shift == 0: continue

            day_type = get_day_type(day)
            if not day_type: continue

            for t, shifts in shifts_info[day_type].items():
                for s in shifts.keys():
                    if allowed_shift == -1 or allowed_shift == s:
                        work[(d, day, t, s)] = model.NewBoolVar(f'w_{d}_d{day}_t{t}_s{s}')

    print("Шаг 2/5: Настройка табу на дыры")
    for day in range(1, num_days + 1):
        day_type = get_day_type(day)
        if not day_type: continue
        for t, shifts in shifts_info[day_type].items():
            for s in shifts.keys():
                possible_drivers = [work[(d, day, t, s)] for d in drivers if (d, day, t, s) in work]
                if possible_drivers:
                    model.AddExactlyOne(possible_drivers)

    print("Шаг 3/5: Запрет на работу >1 смены в день")
    for d in drivers:
        for day in range(1, num_days + 1):
            shifts_today = [work[(d, day, t, s)] for t in shifts_info.get(get_day_type(day), {}) for s in
                            shifts_info.get(get_day_type(day), {}).get(t, {}) if (d, day, t, s) in work]
            if shifts_today:
                model.AddAtMostOne(shifts_today)

    print("Шаг 4/5: Матрица отдыха")
    for d in drivers:
        working_days = [day for day in range(1, num_days + 1) if availabilities[d].get(day, 0) != 0]
        for i in range(len(working_days) - 1):
            day_today = working_days[i]
            day_next = working_days[i + 1]
            type_today = get_day_type(day_today)
            type_next = get_day_type(day_next)
            if not type_today or not type_next: continue

            for t_today, shifts_today in shifts_info[type_today].items():
                for s_today, info_today in shifts_today.items():
                    if (d, day_today, t_today, s_today) not in work: continue

                    end_mins_today = (day_today - 1) * 1440 + info_today["end"]
                    req_rest_mins = (info_today["duration"] * 2) + (2520 if day_next > day_today + 1 else 0)

                    conflicting_shifts_next_day = []

                    for t_next, shifts_next in shifts_info[type_next].items():
                        for s_next, info_next in shifts_next.items():
                            if (d, day_next, t_next, s_next) not in work: continue

                            start_mins_next = (day_next - 1) * 1440 + info_next["start"]
                            if start_mins_next - end_mins_today < req_rest_mins:
                                conflicting_shifts_next_day.append(work[(d, day_next, t_next, s_next)])

                    if conflicting_shifts_next_day:
                        model.Add(sum(conflicting_shifts_next_day) == 0).OnlyEnforceIf(
                            work[(d, day_today, t_today, s_today)])

    print("Шаг 5/5: Настройка целевой функции")
    is_used = {}
    for i, d in enumerate(drivers):
        is_used[d] = model.NewBoolVar(f'used_{d}')
        all_shifts_d = [work[(d, day, t, s)] for day in range(1, num_days + 1) for t in
                        shifts_info.get(get_day_type(day), {}) for s in
                        shifts_info.get(get_day_type(day), {}).get(t, {}) if (d, day, t, s) in work]
        if all_shifts_d:
            model.AddMaxEquality(is_used[d], all_shifts_d)
        else:
            model.Add(is_used[d] == 0)

    total_drivers_used = sum(is_used.values())
    symmetry_breaker = sum(is_used[d] * i for i, d in enumerate(drivers))
    model.Minimize((total_drivers_used * 10_000) + symmetry_breaker)

    print(f"\nЗАПУСК РЕШАТЕЛЯ (Выделено {MAX_SOLVE_TIME} секунд)...")
    solver = cp_model.CpSolver()

    solver.parameters.num_search_workers = 8
    solver.parameters.log_search_progress = True
    solver.parameters.max_time_in_seconds = MAX_SOLVE_TIME

    status = solver.Solve(model)

    if status == cp_model.OPTIMAL or status == cp_model.FEASIBLE:
        used_drivers = set()
        driver_timeline = {d: [] for d in drivers}

        for day in range(1, num_days + 1):
            day_type = get_day_type(day)
            if not day_type: continue
            for t, shifts in shifts_info[day_type].items():
                for s, info in shifts.items():
                    for d in drivers:
                        if (d, day, t, s) in work and solver.Value(work[(d, day, t, s)]) == 1:
                            used_drivers.add(d)
                            abs_start = (day - 1) * 1440 + info["start"]
                            abs_end = (day - 1) * 1440 + info["end"]
                            driver_timeline[d].append(
                                {"day": day, "shift": s, "tram": t, "abs_start": abs_start, "abs_end": abs_end,
                                 "duration": info["duration"]})

        for d in drivers:
            driver_timeline[d].sort(key=lambda x: x["abs_start"])
            shifts_list = driver_timeline[d]
            for i in range(len(shifts_list)):
                if i < len(shifts_list) - 1:
                    shifts_list[i]["rest_after"] = shifts_list[i + 1]["abs_start"] - shifts_list[i]["abs_end"]
                else:
                    shifts_list[i]["rest_after"] = None

        print("\n" + "=" * 50)
        print("РАСПРЕДЕЛЕНИЕ УСПЕШНО ЗАВЕРШЕНО!")
        print(f"Всего закрыто смен: {sum(1 for d in used_drivers for s in driver_timeline[d])}")
        print(f"Использовано водителей: {len(used_drivers)}")
        print("=" * 50)

        export_drivers_report(drivers, used_drivers, driver_timeline, shifts_info, availabilities, get_day_type,
                              driver_schedules, num_days)
        export_schedule_book(driver_timeline, shifts_info, get_day_type, year, month, num_days)

    else:
        print("\nРЕШЕНИЕ НЕ НАЙДЕНО (INFEASIBLE).")
        print("Правило 42 часов заблокировало возможность составить расписание для текущих водителей.")


if __name__ == "__main__":
    solve_schedule()