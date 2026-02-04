import pandas as pd
import os
import calendar
from datetime import date
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from src.logger import get_logger

logger = get_logger(__name__)

SHIFT_1 = 1
SHIFT_2 = 2
REST = 'В'

thin = Side(border_style="thin", color="000000")
medium = Side(border_style="medium", color="000000")

# маски циклов
CYCLIC_PATTERNS = {
    '4x2': [1, 1, 1, 1, 0, 0],
    '5x2': [1, 1, 1, 1, 1, 0, 0],
    '3x2x3x1': [1, 1, 1, 0, 0, 1, 1, 1, 0],
}

# Конфигурация файлов
FILES_CONFIG = [
    {"file": "Приложение №1 График 4х2.xlsx", "subfolder": "4x2", "pattern": "4x2", "step": 12},
    {"file": "Приложение №2 График 5х2.xlsx", "subfolder": "5x2", "pattern": "5x2", "step": 14},
    {"file": "Приложение №3 График 3х2х3х1.xlsx", "subfolder": "3x2x3x1", "pattern": "3x2x3x1", "step": 9},
]


def normalize_key(s):
    if pd.isna(s): return ""
    return str(s).lower().replace(' ', '').replace('*', 'x').replace('х', 'x')


def build_full_cycle(pattern_name, shift_mode, mask):
    """
    4х2: 1 1 1 1 В В 2 2 2 2 В В
    5x2: 1 1 1 1 1 В В 2 2 2 2 2 В В
    3х2х3х1: 1 1 1 В В 2 2 2 В
    """
    mode = normalize_key(shift_mode)

    if '1x2' in mode:
        full_cycle = []
        current_shift = 1

        # Проходим маску дважды, чтобы гарантированно вернуться к начальной смене
        for _ in range(2):
            i = 0
            while i < len(mask):
                if mask[i] == 1:
                    while i < len(mask) and mask[i] == 1:
                        full_cycle.append(current_shift)
                        i += 1
                    current_shift = 2 if current_shift == 1 else 1
                else:
                    full_cycle.append(REST)
                    i += 1

            if current_shift == 1 and len(full_cycle) >= len(mask):
                break
        return full_cycle

    if '2' in mode and '1' not in mode:
        return [(SHIFT_2 if x else REST) for x in mask]
    return [(SHIFT_1 if x else REST) for x in mask]


def solve_cyclic(feb_vals, pattern_name, shift_mode):
    pkey = normalize_key(pattern_name)
    mask = CYCLIC_PATTERNS.get(pkey)
    if not mask: return None

    cycle = build_full_cycle(pattern_name, shift_mode, mask)
    cycle_len = len(cycle)
    best_offset, max_score = -1, -1

    for offset in range(cycle_len):
        score, mismatch = 0, False
        for i, val in enumerate(feb_vals):
            if i >= 28: break
            actual = str(val).strip().upper()
            if actual in ['B', 'V', 'В', '0', 'NAN', '']: actual = REST
            theo = str(cycle[(offset + i) % cycle_len])

            if (actual in ['1', '2']) != (theo in ['1', '2']):
                mismatch = True
                break
            score += 2 if actual == theo else 1

        if not mismatch and score > max_score:
            max_score, best_offset = score, offset

    if best_offset == -1: return None
    # 1 февраля - это 32-й день года (индекс 31)
    jan1_offset = (best_offset - 31) % cycle_len
    return [cycle[(jan1_offset + i) % cycle_len] for i in range(365)]


def format_excel_file(filepath, step):
    wb = load_workbook(filepath)
    ws = wb.active

    font_main = Font(name='Verdana', size=12)
    align_center = Alignment(horizontal='center', vertical='center')

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.font = font_main
            cell.alignment = align_center
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.border = Border(left=thin, right=thin, top=thin, bottom=medium)

    for r_idx in range(1 + step, ws.max_row + 1, step):
        for c_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=medium)

    for col in ws.columns:
        col_letter = col[0].column_letter
        header_val = str(col[0].value)
        if header_val.isdigit():
            ws.column_dimensions[col_letter].width = 5
        elif header_val in ['График', 'Смена', 'Режим']:
            ws.column_dimensions[col_letter].width = 12
        else:
            ws.column_dimensions[col_letter].width = 16

    wb.save(filepath)


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(os.path.dirname(script_dir))
    base_data_dir = os.path.join(project_root, "env_synthetic", "data")
    output_base = os.path.join(base_data_dir, "tabeles_2026_generated")

    for config in FILES_CONFIG:
        input_path = os.path.join(base_data_dir, config["file"])
        if not os.path.exists(input_path):
            logger.error(f"Файл не найден: {input_path}")
            continue

        logger.info(f"Генерация и оформление табеля: {config['subfolder']}...")
        df = pd.read_excel(input_path, dtype=str)
        df.columns = [str(c).strip() for c in df.columns]

        day_cols = [str(i) for i in range(1, 29) if str(i) in df.columns]
        meta_cols = [c for c in ['Таб. №', 'Таб.№', 'График', 'Смена', 'Режим'] if c in df.columns]

        results = {}
        for idx, row in df.iterrows():
            mode = row.get('Смена', row.get('Режим', '1x2'))
            grafik = row.get('График', config['pattern'])
            results[idx] = solve_cyclic([row.get(d, '') for d in day_cols], grafik, mode) or [REST] * 365

        target_dir = os.path.join(output_base, config["subfolder"])
        os.makedirs(target_dir, exist_ok=True)

        for month in range(1, 13):
            days_in_month = calendar.monthrange(2026, month)[1]
            start_doy = date(2026, month, 1).timetuple().tm_yday - 1

            month_df = df[meta_cols].copy()
            for d in range(1, days_in_month + 1):
                month_df[str(d)] = [results[idx][start_doy + d - 1] for idx in df.index]

            month_name = calendar.month_name[month].lower()
            save_path = os.path.join(target_dir, f"{month:02d}_{month_name}_2026.xlsx")
            month_df.to_excel(save_path, index=False)
            format_excel_file(save_path, config['step'])

    logger.info("\nГотово! Все табели продлены.")


if __name__ == "__main__":
    main()
